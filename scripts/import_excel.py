"""Importe les données du classeur GESTION_DES_POINTAGE.xlsx dans la base de l'API.

    python scripts/import_excel.py GESTION_DES_POINTAGE.xlsx --reset
    # classeur protégé par mot de passe : --password ... ou variable POINTAGE_XLSX_PASSWORD (sinon saisie masquée)
"""
import argparse
import datetime as dt
import getpass
import io
import os
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import msoffcrypto  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.database import Base, SessionLocal, engine  # noqa: E402
from app.models import (Affectation, Agent, Autorisation, JourFerie, Pointage, Service)  # noqa: E402
from app.schemas import Motif  # noqa: E402


def ouvrir_classeur(chemin: str, mot_de_passe: str | None = None):
    """Ouvre le classeur (valeurs en cache), en le déchiffrant si nécessaire."""
    if zipfile.is_zipfile(chemin):
        return load_workbook(chemin, data_only=True)
    mdp = mot_de_passe or os.environ.get("POINTAGE_XLSX_PASSWORD") or getpass.getpass("Mot de passe du classeur : ")
    tampon = io.BytesIO()
    with open(chemin, "rb") as f:
        fichier = msoffcrypto.OfficeFile(f)
        fichier.load_key(password=mdp)
        fichier.decrypt(tampon)
    tampon.seek(0)
    return load_workbook(tampon, data_only=True)


def _norm_entete(h) -> str:
    return " ".join(str(h or "").split()).upper()


def lire_tableau(wb, nom_table: str) -> list[dict]:
    for ws in wb.worksheets:
        if nom_table in ws.tables:
            lignes = list(ws[ws.tables[nom_table].ref])
            entetes = [_norm_entete(c.value) for c in lignes[0]]
            return [dict(zip(entetes, (c.value for c in l))) for l in lignes[1:]]
    raise KeyError(nom_table)


def _date(v):
    return v.date() if isinstance(v, dt.datetime) else v if isinstance(v, dt.date) else None


def _heure(v):
    return v.time() if isinstance(v, dt.datetime) else v if isinstance(v, dt.time) else None


def _txt(v) -> str | None:
    return " ".join(str(v).split()) if v not in (None, "") else None


def importer(wb, db) -> dict:
    n = dict(agents=0, affectations=0, pointages=0, doublons_pointages=0, autorisations=0,
             jours_feries=0, motifs_inconnus=0, lignes_ignorees=0)

    elem = wb["ELEMENT"]
    for r in range(2, 20):
        code = _txt(elem.cell(r, 1).value)
        if code and not db.get(Service, code.upper()):
            db.add(Service(code=code.upper()))
    db.flush()

    for l in lire_tableau(wb, "Tableau1"):
        mat, svc, prise = _txt(l.get("N° MATRICULE")), _txt(l.get("SERVICE")), _date(l.get("DATE DE PRISE DE SERVICE"))
        if not (mat and svc and prise):
            n["lignes_ignorees"] += 1
            continue
        mat, svc = mat.upper(), svc.upper()
        if not db.get(Service, svc):
            db.add(Service(code=svc))
        if not db.get(Agent, mat):
            db.add(Agent(matricule=mat, prenom=_txt(l.get("PRENOM")) or "", nom=_txt(l.get("NOM")) or "",
                         sexe=(_txt(l.get("SEXE")) or "HOMME").upper(), fonction=_txt(l.get("FONCTION"))))
            n["agents"] += 1
        db.flush()
        etat = "AFFECTE" if (_txt(l.get("ETATS")) or "").upper().startswith("AFFECTE") else "PRESENT"
        db.add(Affectation(matricule=mat, service_code=svc, etat=etat, date_prise_service=prise,
                           date_affectation=_date(l.get("DATE AFFECTATION"))))
        n["affectations"] += 1
    db.flush()

    for l in lire_tableau(wb, "Tableau3"):
        d, nom = _date(l.get("DATE")), _txt(l.get("NOM DE LA FETE"))
        if d and nom and not db.get(JourFerie, d):
            db.add(JourFerie(date=d, nom=nom))
            n["jours_feries"] += 1
    db.flush()

    vus = set()
    for l in lire_tableau(wb, "Tableau2"):
        d, mat, svc = _date(l.get("DATE")), _txt(l.get("N° MATRICULE")), _txt(l.get("SERVICE"))
        if not (d and mat and svc):
            n["lignes_ignorees"] += 1
            continue
        cle = (d, mat.upper(), svc.upper())
        if cle in vus:
            n["doublons_pointages"] += 1
            continue
        vus.add(cle)
        db.add(Pointage(date=d, heure_arrivee=_heure(l.get("HEURE D'ARRIVE")), heure_depart=_heure(l.get("HEURE DEPART")),
                        matricule=cle[1], service_code=cle[2]))
        n["pointages"] += 1

    motifs = {m.value for m in Motif}
    for l in lire_tableau(wb, "Tableau4"):
        mat, svc = _txt(l.get("N° MATRICULE")), _txt(l.get("SERVICE"))
        dep, ret = _date(l.get("DATE DE DEPART")), _date(l.get("DATE DE RETOUR"))
        if not (mat and svc and dep and ret):
            n["lignes_ignorees"] += 1
            continue
        motif = (_txt(l.get("MOTIF OU JUSTIFICATION D'ABSENCE")) or "AUTRE").upper()
        if motif not in motifs:
            motif, n["motifs_inconnus"] = "AUTRE", n["motifs_inconnus"] + 1
        annule = (_txt(l.get("ANNULATION")) or "").upper().startswith("ANNUL")
        db.add(Autorisation(matricule=mat.upper(), service_code=svc.upper(), date_depart=dep, date_retour=ret,
                            motif=motif, annule=annule))
        n["autorisations"] += 1
    db.commit()
    return n


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("classeur")
    ap.add_argument("--password", help="mot de passe du classeur (sinon variable POINTAGE_XLSX_PASSWORD ou saisie)")
    ap.add_argument("--reset", action="store_true", help="vider la base avant l'import")
    a = ap.parse_args()

    if a.reset:
        Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        if db.query(Agent).count():
            sys.exit("La base contient déjà des agents. Relancer avec --reset pour repartir de zéro.")
        print(importer(ouvrir_classeur(a.classeur, a.password), db))


if __name__ == "__main__":
    main()
